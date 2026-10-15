import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
p=sys.argv[1]; d=Document(p)
for st in d.styles:
    try: f=st.font
    except Exception: continue
    if f is None: continue
    f.name='Malgun Gothic'
    rpr=st.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),'Malgun Gothic')
for nm,sz in [('Normal',10),('Body Text',10),('First Paragraph',10),('Compact',9.5),('Heading 1',15),('Heading 2',12.5),('Title',22),('Subtitle',12),('Source Code',8.5)]:
    try: d.styles[nm].font.size=Pt(sz)
    except KeyError: pass
for nm in ['Heading 1','Heading 2','Title']:
    try: d.styles[nm].font.color.rgb=RGBColor(0x1f,0x3a,0x5f)
    except KeyError: pass
sec=d.sections[0]; sec.left_margin=sec.right_margin=Cm(2); sec.top_margin=sec.bottom_margin=Cm(2)
TW=17.0
def wlen(s): return sum(2 if ord(ch)>0x1100 else 1 for ch in s)
for t in d.tables:
    tbl=t._tbl; pr=tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'B7C0CC'); b.append(el)
    pr.append(b)
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); pr.append(lay)
    ncol=len(t.columns)
    lens=[0]*ncol
    for row in t.rows:
        for i,c in enumerate(row.cells[:ncol]):
            lens[i]=max(lens[i], min(60, wlen(c.text)))
    lens=[max(9,l) for l in lens]; tot=sum(lens)
    widths=[Cm(TW*l/tot) for l in lens]
    grid=tbl.tblGrid
    for i,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[i].twips)))
    for row in t.rows:
        for i,c in enumerate(row.cells[:ncol]):
            c.width=widths[i]
            for para in c.paragraphs:
                for r in para.runs: r.font.size=Pt(8.5)
    for c in t.rows[0].cells:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),'E8EEF5'); tcPr.append(sh)
        for para in c.paragraphs:
            for r in para.runs: r.bold=True
d.save(p)
